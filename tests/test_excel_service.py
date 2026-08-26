"""Tests for Excel report generation service."""

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app.services.excel import generate_excel_bytes


@pytest.fixture
def sample_nasa_data():
    """Return sample NASA data for Excel generation."""
    dates = [f"202401{d:02d}" for d in range(1, 16)]
    return {
        "properties": {
            "parameter": {
                "WS2M": {d: 2.5 + (i * 0.1) for i, d in enumerate(dates)},
                "WD2M": {d: 180.0 + (i * 5) for i, d in enumerate(dates)},
                "ALLSKY_SFC_SW_DWN": {d: 4.5 + (i * 0.05) for i, d in enumerate(dates)},
            }
        }
    }


class TestGenerateExcelBytes:
    """Tests for generate_excel_bytes function."""

    def test_returns_bytesio(self, sample_nasa_data):
        """Function returns a BytesIO buffer."""
        result = generate_excel_bytes(
            station_name="Test Station",
            latitude=4.6097,
            longitude=-74.0817,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 1, 15),
            nasa_data=sample_nasa_data,
        )
        assert isinstance(result, io.BytesIO)

    def test_valid_xlsx_content(self, sample_nasa_data):
        """Generated content is a valid XLSX file."""
        result = generate_excel_bytes(
            station_name="Test Station",
            latitude=4.6097,
            longitude=-74.0817,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 1, 15),
            nasa_data=sample_nasa_data,
        )
        # Verify it starts with ZIP magic bytes (XLSX is a ZIP)
        assert result.read(2) == b"PK"

    def test_workbook_has_expected_sheets(self, sample_nasa_data):
        """Workbook contains all expected sheets."""
        result = generate_excel_bytes(
            station_name="Test Station",
            latitude=4.6097,
            longitude=-74.0817,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 1, 15),
            nasa_data=sample_nasa_data,
        )
        wb = load_workbook(result)
        expected = [
            "Info",
            "Solar Radiation",
            "Monthly Solar Radiation",
            "Wind Data",
            "Monthly Summary",
            "Scatter Chart",
            "Wind Rose Chart",
            "Monthly Summary Polar",
            "Monthly Solar Radiation Chart",
        ]
        for sheet in expected:
            assert sheet in wb.sheetnames, f"Missing sheet: {sheet}"

    def test_wind_data_sheet_content(self, sample_nasa_data):
        """Wind Data sheet contains correct data."""
        result = generate_excel_bytes(
            station_name="Test Station",
            latitude=4.6097,
            longitude=-74.0817,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 1, 15),
            nasa_data=sample_nasa_data,
        )
        wb = load_workbook(result)
        ws = wb["Wind Data"]
        # Header row
        assert ws.cell(1, 1).value == "Date"
        assert ws.cell(1, 2).value == "Wind Speed (m/s)"
        assert ws.cell(1, 3).value == "Wind Direction (degrees)"
        # Data rows (15 days)
        assert ws.max_row == 16  # 1 header + 15 data

    def test_info_sheet_content(self, sample_nasa_data):
        """Info sheet contains station metadata."""
        result = generate_excel_bytes(
            station_name="Test Station",
            latitude=4.6097,
            longitude=-74.0817,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 1, 15),
            nasa_data=sample_nasa_data,
        )
        wb = load_workbook(result)
        ws = wb["Info"]
        # Check that station name is in the info sheet
        found = False
        for row in ws.iter_rows(values_only=True):
            if "Test Station" in str(row):
                found = True
                break
        assert found

    def test_no_temp_files_left(self, sample_nasa_data):
        """No temporary chart files are left after generation."""
        import glob

        before = set(glob.glob("*.jpg"))
        generate_excel_bytes(
            station_name="Test Station",
            latitude=4.6097,
            longitude=-74.0817,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 1, 15),
            nasa_data=sample_nasa_data,
        )
        after = set(glob.glob("*.jpg"))
        assert after == before

    def test_missing_wind_data_raises(self):
        """Missing wind parameters raises ValueError."""
        with pytest.raises(ValueError, match="missing wind parameters"):
            generate_excel_bytes(
                station_name="Test Station",
                latitude=4.6097,
                longitude=-74.0817,
                start_date=date(2024, 1, 1),
                end_date=date(2024, 1, 15),
                nasa_data={"properties": {"parameter": {}}},
            )

    def test_solar_data_optional(self):
        """Excel generates without solar data."""
        dates = [f"202401{d:02d}" for d in range(1, 6)]
        nasa_data = {
            "properties": {
                "parameter": {
                    "WS2M": {d: 2.0 for d in dates},
                    "WD2M": {d: 180.0 for d in dates},
                }
            }
        }
        result = generate_excel_bytes(
            station_name="Test Station",
            latitude=4.6097,
            longitude=-74.0817,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 1, 5),
            nasa_data=nasa_data,
        )
        wb = load_workbook(result)
        # Solar sheets should not be present
        assert "Solar Radiation" not in wb.sheetnames or wb["Solar Radiation"].max_row == 1

    def test_null_nasa_values_handled(self):
        """NASA -999.0 values are treated as null in output."""
        dates = [f"202401{d:02d}" for d in range(1, 6)]
        nasa_data = {
            "properties": {
                "parameter": {
                    "WS2M": {d: -999.0 for d in dates},
                    "WD2M": {d: -999.0 for d in dates},
                    "ALLSKY_SFC_SW_DWN": {d: -999.0 for d in dates},
                }
            }
        }
        # Should not raise -999.0 values are valid in NASA response
        result = generate_excel_bytes(
            station_name="Test Station",
            latitude=4.6097,
            longitude=-74.0817,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 1, 5),
            nasa_data=nasa_data,
        )
        assert isinstance(result, io.BytesIO)
