"""Tests for weather API endpoints (integration tests with mocked NASA)."""

import httpx
import respx


class TestWeatherDataEndpoint:
    """Tests for POST /api/v1/weather-data."""

    @respx.mock
    def test_weather_data_success(
        self, client, valid_request, mock_nasa_wind_response, mock_nasa_solar_response
    ):
        """Successful request returns 200 with daily_data, monthly_summary, metadata."""
        respx.get("https://power.larc.nasa.gov/api/temporal/daily/point").mock(
            side_effect=[
                httpx.Response(200, json=mock_nasa_wind_response),
                httpx.Response(200, json=mock_nasa_solar_response),
            ]
        )

        response = client.post("/api/v1/weather-data", json=valid_request)

        assert response.status_code == 200
        data = response.json()
        assert "daily_data" in data
        assert "monthly_summary" in data
        assert "metadata" in data
        assert data["metadata"]["station_name"] == "Bogota Station"
        assert data["metadata"]["latitude"] == 4.6097
        assert data["metadata"]["longitude"] == -74.0817
        assert len(data["daily_data"]) == 31
        assert len(data["monthly_summary"]) >= 1

    @respx.mock
    def test_weather_data_daily_fields(
        self, client, valid_request, mock_nasa_wind_response, mock_nasa_solar_response
    ):
        """Daily data contains expected fields."""
        respx.get("https://power.larc.nasa.gov/api/temporal/daily/point").mock(
            side_effect=[
                httpx.Response(200, json=mock_nasa_wind_response),
                httpx.Response(200, json=mock_nasa_solar_response),
            ]
        )

        response = client.post("/api/v1/weather-data", json=valid_request)
        data = response.json()

        daily = data["daily_data"][0]
        assert "date" in daily
        assert "wind_speed_ms" in daily
        assert "wind_direction_deg" in daily
        assert "solar_radiation_kwh" in daily

    @respx.mock
    def test_weather_data_monthly_fields(
        self, client, valid_request, mock_nasa_wind_response, mock_nasa_solar_response
    ):
        """Monthly summary contains expected fields."""
        respx.get("https://power.larc.nasa.gov/api/temporal/daily/point").mock(
            side_effect=[
                httpx.Response(200, json=mock_nasa_wind_response),
                httpx.Response(200, json=mock_nasa_solar_response),
            ]
        )

        response = client.post("/api/v1/weather-data", json=valid_request)
        data = response.json()

        monthly = data["monthly_summary"][0]
        assert "year_month" in monthly
        assert "avg_wind_speed_ms" in monthly
        assert "avg_wind_direction_deg" in monthly
        assert "avg_solar_radiation_kwh" in monthly

    @respx.mock
    def test_weather_data_metadata_fields(
        self, client, valid_request, mock_nasa_wind_response, mock_nasa_solar_response
    ):
        """Metadata contains expected fields."""
        respx.get("https://power.larc.nasa.gov/api/temporal/daily/point").mock(
            side_effect=[
                httpx.Response(200, json=mock_nasa_wind_response),
                httpx.Response(200, json=mock_nasa_solar_response),
            ]
        )

        response = client.post("/api/v1/weather-data", json=valid_request)
        data = response.json()

        meta = data["metadata"]
        assert "station_name" in meta
        assert "latitude" in meta
        assert "longitude" in meta
        assert "start_date" in meta
        assert "end_date" in meta
        assert "total_days" in meta

    def test_weather_data_invalid_body(self, client):
        """Empty body returns 422."""
        response = client.post("/api/v1/weather-data", json={})
        assert response.status_code == 422

    def test_weather_data_missing_station_name(self, client):
        """Missing station_name returns 422."""
        response = client.post(
            "/api/v1/weather-data",
            json={
                "latitude": 4.6097,
                "longitude": -74.0817,
                "start": "2024-01-01",
                "end": "2024-01-31",
            },
        )
        assert response.status_code == 422

    def test_weather_data_empty_station_name(self, client):
        """Empty station_name returns 422."""
        response = client.post(
            "/api/v1/weather-data",
            json={
                "station_name": "",
                "latitude": 4.6097,
                "longitude": -74.0817,
                "start": "2024-01-01",
                "end": "2024-01-31",
            },
        )
        assert response.status_code == 422

    def test_weather_data_invalid_coordinates(self, client):
        """Coordinates outside Colombia return 422."""
        response = client.post(
            "/api/v1/weather-data",
            json={
                "station_name": "Test",
                "latitude": 50.0,
                "longitude": -74.0,
                "start": "2024-01-01",
                "end": "2024-01-31",
            },
        )
        assert response.status_code == 422

    def test_weather_data_invalid_dates(self, client):
        """Start after end returns 422."""
        response = client.post(
            "/api/v1/weather-data",
            json={
                "station_name": "Test",
                "latitude": 4.6097,
                "longitude": -74.0817,
                "start": "2024-01-31",
                "end": "2024-01-01",
            },
        )
        assert response.status_code == 422

    @respx.mock
    def test_weather_data_nasa_error_returns_502(
        self, client, valid_request
    ):
        """NASA API error returns 502."""
        respx.get("https://power.larc.nasa.gov/api/temporal/daily/point").mock(
            return_value=httpx.Response(500, text="Internal Server Error")
        )

        response = client.post("/api/v1/weather-data", json=valid_request)
        assert response.status_code == 502


class TestExcelReportEndpoint:
    """Tests for POST /api/v1/excel-report."""

    @respx.mock
    def test_excel_report_success(
        self, client, valid_request, mock_nasa_wind_response, mock_nasa_solar_response
    ):
        """Successful request returns streaming Excel file."""
        respx.get("https://power.larc.nasa.gov/api/temporal/daily/point").mock(
            side_effect=[
                httpx.Response(200, json=mock_nasa_wind_response),
                httpx.Response(200, json=mock_nasa_solar_response),
            ]
        )

        response = client.post("/api/v1/excel-report", json=valid_request)

        assert response.status_code == 200
        assert (
            response.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "attachment" in response.headers.get("content-disposition", "")
        # Verify it's a valid xlsx (starts with PK zip header)
        assert response.content[:2] == b"PK"

    @respx.mock
    def test_excel_report_is_valid_workbook(
        self, client, valid_request, mock_nasa_wind_response, mock_nasa_solar_response
    ):
        """Generated Excel is a valid workbook with expected sheets."""
        import io

        from openpyxl import load_workbook

        respx.get("https://power.larc.nasa.gov/api/temporal/daily/point").mock(
            side_effect=[
                httpx.Response(200, json=mock_nasa_wind_response),
                httpx.Response(200, json=mock_nasa_solar_response),
            ]
        )

        response = client.post("/api/v1/excel-report", json=valid_request)
        wb = load_workbook(io.BytesIO(response.content))

        expected_sheets = [
            "Info",
            "Wind Data",
            "Monthly Summary",
            "Scatter Chart",
            "Wind Rose Chart",
        ]
        for sheet_name in expected_sheets:
            assert sheet_name in wb.sheetnames, f"Missing sheet: {sheet_name}"

    @respx.mock
    def test_excel_report_no_persistent_files(
        self, client, valid_request, mock_nasa_wind_response, mock_nasa_solar_response, tmp_path
    ):
        """No .xlsx files should be left in the working directory."""
        import glob

        respx.get("https://power.larc.nasa.gov/api/temporal/daily/point").mock(
            side_effect=[
                httpx.Response(200, json=mock_nasa_wind_response),
                httpx.Response(200, json=mock_nasa_solar_response),
            ]
        )

        # Count xlsx files before
        before = set(glob.glob("*.xlsx"))

        client.post("/api/v1/excel-report", json=valid_request)

        # Count xlsx files after
        after = set(glob.glob("*.xlsx"))

        # No new xlsx files should have been created
        assert after == before

    def test_excel_report_invalid_body(self, client):
        """Empty body returns 422."""
        response = client.post("/api/v1/excel-report", json={})
        assert response.status_code == 422


class TestOldEndpointsRemoved:
    """Tests that old insecure endpoints are removed."""

    def test_generate_files_removed(self, client):
        """Old /generate-files endpoint is removed."""
        response = client.post("/generate-files", json={})
        assert response.status_code == 404

    def test_download_removed(self, client):
        """Old /download/{filename} endpoint is removed."""
        response = client.get("/download/test.xlsx")
        assert response.status_code == 404
